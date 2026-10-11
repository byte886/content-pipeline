#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
转写文本近重复普查（near-duplicate scan）。

md5 只能发现"二进制完全相同"的重复（如一稿两发、下载串流）；本脚本用
**字符 bigram 的 TF-IDF 余弦 + bigram 集合 Jaccard** 发现"同一段口播被
掐头去尾/重新剪辑/分段重发"等轻度重复，无需分词、纯标准库实现。

用法：
  python3 scripts/transcript_dup_scan.py <transcripts目录> [--xlsx 台账.xlsx] \
      [--threshold 0.4] [--json out.json]

约定：转写目录结构为 <transcripts>/<视频名>/transcript.md；
脚本自动跳过 "# 标题"、"## 第N段"、"> 元信息" 行，只比对口播正文，
并剔除标点/空白/符号，仅保留中文与字母数字。
"""
import argparse
import glob
import json
import math
import os
import re
from collections import Counter


def read_body(md_path):
    body = []
    with open(md_path, encoding="utf-8") as f:
        for ln in f.read().splitlines():
            s = ln.strip()
            if not s or s.startswith("#") or s.startswith(">"):
                continue
            body.append(s)
    txt = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fa5%]", "", "".join(body))
    return txt


def bigrams(t):
    return [t[i:i + 2] for i in range(len(t) - 1)]


def load_xlsx_map(xlsx_path, transcripts_dir):
    """local_path(basename 去 .mp4) -> {vid,row,title}。xlsx 缺失则返回空。"""
    m = {}
    if not xlsx_path or not os.path.exists(xlsx_path):
        return m
    try:
        import openpyxl
    except ImportError:
        return m
    ws = openpyxl.load_workbook(xlsx_path, read_only=True).active
    for r in range(2, ws.max_row + 1):
        url = ws.cell(r, 1).value or ""
        title = ws.cell(r, 2).value or ""
        local = ws.cell(r, 9).value or ""
        vid = ""
        mm = re.search(r"/video/(\d+)", str(url))
        if mm:
            vid = mm.group(1)
        stem = os.path.splitext(os.path.basename(str(local)))[0]
        if stem:
            m[stem] = {"vid": vid, "row": r, "title": str(title)}
    return m


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("transcripts")
    ap.add_argument("--xlsx", default="")
    ap.add_argument("--threshold", type=float, default=0.4)
    ap.add_argument("--json", default="")
    args = ap.parse_args()

    docs = []
    for md in sorted(glob.glob(os.path.join(args.transcripts, "*", "transcript.md"))):
        stem = os.path.basename(os.path.dirname(md))
        body = read_body(md)
        docs.append({"stem": stem, "body": body, "n": len(body),
                     "tf": Counter(bigrams(body))})
    n_docs = len(docs)

    # document frequency -> idf
    df = Counter()
    for d in docs:
        for g in d["tf"].keys():
            df[g] += 1
    idf = {g: math.log((1 + n_docs) / (1 + c)) + 1.0 for g, c in df.items()}

    # normalized tf-idf sparse vectors
    for d in docs:
        vec = {}
        norm = 0.0
        for g, c in d["tf"].items():
            w = (1 + math.log(c)) * idf[g]      # sublinear tf
            vec[g] = w
            norm += w * w
        norm = math.sqrt(norm) or 1.0
        d["vec"] = {g: w / norm for g, w in vec.items()}
        d["set"] = set(d["tf"].keys())

    xmap = load_xlsx_map(args.xlsx, args.transcripts)
    pairs = []
    for i in range(n_docs):
        vi = docs[i]["vec"]
        for j in range(i + 1, n_docs):
            vj = docs[j]["vec"]
            small, large = (vi, vj) if len(vi) <= len(vj) else (vj, vi)
            cos = sum(w * large.get(g, 0.0) for g, w in small.items())
            if cos < args.threshold:
                continue
            A, B = docs[i]["set"], docs[j]["set"]
            jac = len(A & B) / len(A | B) if (A | B) else 0.0
            pairs.append({"cos": round(cos, 4), "jac": round(jac, 4),
                          "a": docs[i]["stem"], "b": docs[j]["stem"],
                          "na": docs[i]["n"], "nb": docs[j]["n"],
                          "ha": docs[i]["body"][:42], "hb": docs[j]["body"][:42]})

    pairs.sort(key=lambda p: p["cos"], reverse=True)
    for p in pairs:
        p["vid_a"] = xmap.get(p["a"], {}).get("vid", "")
        p["vid_b"] = xmap.get(p["b"], {}).get("vid", "")
        p["row_a"] = xmap.get(p["a"], {}).get("row", "")
        p["row_b"] = xmap.get(p["b"], {}).get("row", "")

    print(f"扫描转写文档数：{n_docs}")
    print(f"余弦≥{args.threshold} 的对数：{len(pairs)}")
    for cut in (0.90, 0.80, 0.60, 0.45):
        print(f"  cos≥{cut}: {sum(1 for p in pairs if p['cos'] >= cut)} 对")
    print("=" * 88)
    for k, p in enumerate(pairs, 1):
        tag = ("高度疑似同源" if p["cos"] >= 0.80 else
               "疑似同源/剪辑" if p["cos"] >= 0.60 else "主题相近(待研判)")
        print(f"[{k}] cos={p['cos']} jac={p['jac']}  {tag}")
        print(f"    A 行{p['row_a']} vid={p['vid_a']} ({p['na']}字) {p['a']}")
        print(f"       开头: {p['ha']}")
        print(f"    B 行{p['row_b']} vid={p['vid_b']} ({p['nb']}字) {p['b']}")
        print(f"       开头: {p['hb']}")
    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump({"n_docs": n_docs, "pairs": pairs}, f,
                      ensure_ascii=False, indent=2)
        print(f"\n已写出 JSON：{args.json}")


if __name__ == "__main__":
    main()
