#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
视频画面近重复普查（visual near-duplicate scan）。

文本转写普查（transcript_dup_scan.py）只能发现"口播文本雷同"；本脚本从
**画面**维度发现同源视频，并区分两种情况：
  - 同一视频文件重发 / 剪辑 / 切片：出现多张逐帧几乎相同的画面（距离≈0、占比≈1）；
  - 同稿重新表演（重拍）：文本雷同，但没有逐帧相同画面。

方法（零第三方依赖，只需 ffmpeg/ffprobe）：
  每个 mp4 用快速 seek 均匀抽 N 帧，ffmpeg 直接降采样成 (gw+1)xgh 灰度小图，
  纯 Python 计算 dHash（gw x gh 相邻像素差分）；两视频间对每帧找最近 Hamming
  距离，统计"近完全相同帧"（距离 ≤ --dist）数量与占比。

重要经验（郭颖库 2026-10-12 实测）：
  同一作者常在**固定演播室、固定机位、同一场连讲多条**（如钻石系列、黄金
  系列、瑞兽系列），9x8 粗 dHash 会把这些"同场不同内容"误报成大量高相似对
  （73 对，跨主题连通）。因此：
    1. 默认用 16x16（256 bit）细 dHash，能区分手势/字幕/道具/插图差异；
    2. 不看整体均值，只看是否存在多张距离极小的"几乎同一张图"；
    3. 判重需满足强条件：最小距离≈0 且近同帧占比≥0.8（同文件/同片）；
       占比 0.25–0.8 且文本不同的，基本是同场连讲，须抽帧肉眼复核，不能直接判重。

用法：
  python3 scripts/video_visual_dup_scan.py <视频目录> [--xlsx 台账.xlsx] \
      [--frames 8] [--gw 16] [--gh 16] [--dist 0(自动)] \
      [--strong-ratio 0.8] [--json out.json]
"""
import argparse
import glob
import json
import os
import re
import subprocess


def duration_of(path):
    r = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=nk=1:nw=1", path],
        capture_output=True, text=True)
    try:
        return float(r.stdout.strip())
    except ValueError:
        return 0.0


def dhash(b, gw, gh):
    """(gw+1)*gh 字节灰度图；每行相邻 gw 次比较 -> gw*gh bit。"""
    h = 0
    bit = 0
    rowbytes = gw + 1
    for row in range(gh):
        base = row * rowbytes
        for c in range(gw):
            if b[base + c] > b[base + c + 1]:
                h |= 1 << bit
            bit += 1
    return h


def frame_hashes(path, n, gw, gh):
    dur = duration_of(path)
    if dur <= 0:
        return [], 0.0
    raw = (gw + 1) * gh
    hs = []
    for k in range(n):
        t = (k + 0.5) * dur / n
        r = subprocess.run(
            ["ffmpeg", "-nostdin", "-loglevel", "error",
             "-ss", f"{t:.3f}", "-i", path, "-frames:v", "1",
             "-vf", f"scale={gw + 1}:{gh},format=gray", "-f", "rawvideo", "-"],
            capture_output=True)
        b = r.stdout
        if len(b) >= raw:
            hs.append(dhash(b[:raw], gw, gh))
    return hs, dur


def ham(a, b):
    return (a ^ b).bit_count()


def load_xlsx_map(xlsx_path):
    m = {}
    if not xlsx_path or not os.path.exists(xlsx_path):
        return m
    try:
        import openpyxl
    except ImportError:
        return m
    ws = openpyxl.load_workbook(xlsx_path, read_only=True).active
    for r in range(2, ws.max_row + 1):
        url = str(ws.cell(r, 1).value or "")
        local = str(ws.cell(r, 9).value or "")
        mm = re.search(r"/video/(\d+)", url)
        stem = os.path.splitext(os.path.basename(local))[0]
        if stem:
            m[stem] = {"vid": mm.group(1) if mm else "", "row": r}
    return m


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("video_dir")
    ap.add_argument("--xlsx", default="")
    ap.add_argument("--frames", type=int, default=8)
    ap.add_argument("--gw", type=int, default=16, help="dHash 横向格子数")
    ap.add_argument("--gh", type=int, default=16, help="dHash 纵向格子数")
    ap.add_argument("--dist", type=int, default=0,
                    help="近同帧 Hamming 距离阈值；0=自动(约为 bit 数的二十分之一)")
    ap.add_argument("--strong-ratio", type=float, default=0.8,
                    help="近同帧占比≥该值且最小距离很小才判同文件/同片")
    ap.add_argument("--weak-ratio", type=float, default=0.25,
                    help="占比≥该值列为待肉眼复核")
    ap.add_argument("--json", default="")
    args = ap.parse_args()

    bits = args.gw * args.gh
    dist = args.dist if args.dist > 0 else max(4, round(bits * 0.05))

    vids = sorted(glob.glob(os.path.join(args.video_dir, "*.mp4")))
    docs = []
    for i, p in enumerate(vids, 1):
        hs, dur = frame_hashes(p, args.frames, args.gw, args.gh)
        docs.append({"stem": os.path.splitext(os.path.basename(p))[0],
                     "frames": hs, "dur": round(dur, 1)})
        if i % 20 == 0 or i == len(vids):
            print(f"  抽帧 {i}/{len(vids)}", flush=True)

    xmap = load_xlsx_map(args.xlsx)
    weak, top_min = [], []
    for i in range(len(docs)):
        for j in range(i + 1, len(docs)):
            A, B = docs[i]["frames"], docs[j]["frames"]
            if not A or not B:
                continue
            da = [min(ham(x, y) for y in B) for x in A]
            db = [min(ham(y, x) for x in A) for y in B]
            near = (sum(d <= dist for d in da) +
                    sum(d <= dist for d in db)) / 2.0
            ratio = near / min(len(A), len(B))
            mind = min(min(da), min(db))
            rec = {"a": docs[i]["stem"], "b": docs[j]["stem"],
                   "na": len(A), "nb": len(B), "near": round(near, 1),
                   "ratio": round(ratio, 3), "min_dist": mind,
                   "vid_a": xmap.get(docs[i]["stem"], {}).get("vid", ""),
                   "vid_b": xmap.get(docs[j]["stem"], {}).get("vid", ""),
                   "row_a": xmap.get(docs[i]["stem"], {}).get("row", ""),
                   "row_b": xmap.get(docs[j]["stem"], {}).get("row", "")}
            top_min.append(rec)
            if ratio >= args.weak_ratio:
                weak.append(rec)

    weak.sort(key=lambda r: (r["ratio"], r["near"], -r["min_dist"]), reverse=True)
    strong = [r for r in weak if r["ratio"] >= args.strong_ratio and r["min_dist"] <= dist]
    review = [r for r in weak if r not in strong]
    top_min.sort(key=lambda r: r["min_dist"])

    print("=" * 88)
    print(f"视频数 {len(docs)}；dHash {bits}bit；近同帧距离阈值≤{dist}")
    print(f"【强判同文件/同片】占比≥{args.strong_ratio} 且最小距离≤{dist}：{len(strong)} 对")
    for k, r in enumerate(strong, 1):
        print(f"  [S{k}] 近同帧{r['near']:.0f} 占比{r['ratio']:.2f} 最小距离{r['min_dist']} | "
              f"行{r['row_a']} {r['a']}  ↔  行{r['row_b']} {r['b']}")
    print(f"【待肉眼复核】占比 {args.weak_ratio}–{args.strong_ratio}（多为同场连讲，勿直接判重）：{len(review)} 对")
    for k, r in enumerate(review, 1):
        print(f"  [R{k}] 近同帧{r['near']:.0f} 占比{r['ratio']:.2f} 最小距离{r['min_dist']} | "
              f"{r['a'][:20]}  ↔  {r['b'][:20]}")
    print("-" * 88)
    print("全局最小 Hamming 距离 top15（背景相似度基线）：")
    for r in top_min[:15]:
        print(f"  d={r['min_dist']:2d} 近同帧{r['near']:.0f} 占比{r['ratio']:.2f} | "
              f"{r['a'][:20]}  ↔  {r['b'][:20]}")

    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump({"n_videos": len(docs), "bits": bits, "dist": dist,
                       "strong": strong, "review": review, "top_min": top_min[:30]},
                      f, ensure_ascii=False, indent=2)
        print(f"\n已写出 JSON：{args.json}")


if __name__ == "__main__":
    main()
