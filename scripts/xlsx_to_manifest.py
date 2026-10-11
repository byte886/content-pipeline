#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xlsx 视频台账 -> 入库 manifest.csv 同步 / 对账脚本。

定位（权威关系，勿混淆）
  - xlsx 是本机（Mac）富编辑工作副本，被 .gitignore 忽略、不入库、换设备不可见；
  - manifest.csv 是入库、版本化、跨设备 / 跨会话 / 跨仓库可见的「资产权威快照」，
    由本脚本从 xlsx 单向生成；
  - 不要手改 manifest.csv：所有登记 / 评级 / 备注编辑都在 xlsx 完成，提交前重跑本脚本同步。

视频 / 音频本体（*.mp4/*.wav）体积大、继续 gitignore 不入库；manifest 记录其相对路径、
md5 指纹与时长，任何设备凭 url 可重新下载、凭 md5 校验是否同一文件。

用法（在仓库根执行）：
  python3 scripts/xlsx_to_manifest.py \
    --xlsx "library/05_knowledge/jewelry/珠宝学者郭颖/郭颖珠宝视频知识库_总清单.xlsx" \
    --video-dir "library/01_video/jewelry/珠宝学者郭颖" \
    --transcript-dir "library/05_knowledge/jewelry/珠宝学者郭颖/transcripts" \
    --out "library/01_video/jewelry/珠宝学者郭颖/manifest.csv" \
    --platform douyin --account "珠宝学者郭颖"

输出 CSV 固定列：
  vid,url,platform,account,title,category,value_level,review_grade,review_topic,
  downloaded,transcribed,media_relpath,md5,duration_sec,transcript_relpath,review_note

对账：打印各维度计数；出现「已下载但缺 md5 / 已转写但缺文稿 / vid 缺失或重复 /
状态与文件不符」等硬异常时以非零码退出。md5 重复只告警（已知存在同文件一稿两发）。
"""
import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("缺少 openpyxl：pip install openpyxl")

COLUMNS = [
    "vid", "url", "platform", "account", "title", "category", "value_level",
    "review_grade", "review_topic", "downloaded", "transcribed", "media_relpath",
    "md5", "duration_sec", "transcript_relpath", "review_note",
]
# xlsx 表头 -> manifest 字段（其余 platform/account 由参数注入，派生字段另算）
HEADER_MAP = {
    "url": "url", "title": "title_raw", "title_clean": "title",
    "category": "category", "value_level": "value_level",
    "downloaded": "downloaded", "transcribed": "transcribed",
    "local_path": "media_relpath", "review_grade": "review_grade",
    "review_topic": "review_topic", "review_note": "review_note",
}
VID_RE = re.compile(r"\d{15,20}")


def clean(v):
    """单元格 -> 去空白字符串；字段内换行/制表符压成空格，保证 CSV 一行一记录。"""
    if v is None:
        return ""
    s = str(v).strip()
    return re.sub(r"[\r\n\t]+", " ", s)


def extract_vid(url):
    """vid 一律取 url 中的长数字串（xlsx 的 vid 列是科学计数法浮点数，不可用）。"""
    nums = VID_RE.findall(url or "")
    return max(nums, key=len) if nums else ""


def md5_of(path, buf=1 << 20):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(buf), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    ap = argparse.ArgumentParser(description="xlsx 视频台账 -> manifest.csv 同步/对账")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--video-dir", required=True, help="mp4 所在目录（相对仓库根或绝对）")
    ap.add_argument("--transcript-dir", required=True, help="转写 transcripts 目录")
    ap.add_argument("--out", required=True, help="manifest.csv 输出路径")
    ap.add_argument("--platform", default="douyin")
    ap.add_argument("--account", default="")
    args = ap.parse_args()

    repo = Path(__file__).resolve().parents[1]

    def resolve(p):
        pp = Path(p)
        return pp if pp.is_absolute() else repo / pp

    xlsx = resolve(args.xlsx)
    video_dir = resolve(args.video_dir)
    trans_dir = resolve(args.transcript_dir)
    out = resolve(args.out)
    if not xlsx.exists():
        sys.exit(f"找不到 xlsx：{xlsx}")

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    headers = {clean(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}
    missing = [h for h in HEADER_MAP if h not in headers]
    if missing:
        sys.exit(f"xlsx 缺少必需列：{missing}；实际表头：{list(headers)}")

    def cell(row, header):
        return clean(ws.cell(row, headers[header]).value)

    rows = []
    errors, dup_md5 = [], []
    seen_vid = set()
    md5_seen = {}
    for r in range(2, ws.max_row + 1):
        url = cell(r, "url")
        if not url:
            continue
        rec = {k: "" for k in COLUMNS}
        rec["url"] = url
        rec["vid"] = extract_vid(url)
        rec["platform"] = args.platform
        rec["account"] = args.account
        rec["title"] = cell(r, "title_clean") or cell(r, "title")
        rec["category"] = cell(r, "category")
        rec["value_level"] = cell(r, "value_level")
        rec["review_grade"] = cell(r, "review_grade")
        rec["review_topic"] = cell(r, "review_topic")
        rec["downloaded"] = cell(r, "downloaded")
        rec["transcribed"] = cell(r, "transcribed")
        rec["media_relpath"] = cell(r, "local_path")
        rec["review_note"] = cell(r, "review_note")

        is_dl = rec["downloaded"] == "已下载"
        is_tr = rec["transcribed"] == "已转写"

        # vid 完整性 / 唯一性
        if not rec["vid"]:
            errors.append(f"行{r} 无法从 url 提取 vid：{url}")
        elif rec["vid"] in seen_vid:
            errors.append(f"行{r} vid 重复：{rec['vid']}")
        else:
            seen_vid.add(rec["vid"])

        # 媒体文件 + md5
        media = rec["media_relpath"]
        media_path = resolve(media) if media else None
        if is_dl:
            if not media:
                errors.append(f"行{r} 标记已下载但 local_path 为空（{rec['title'][:20]}）")
            elif not media_path.exists():
                errors.append(f"行{r} 已下载但找不到文件：{media}")
            else:
                rec["md5"] = md5_of(media_path)
                if rec["md5"] in md5_seen:
                    dup_md5.append((rec["md5"], md5_seen[rec["md5"]], rec["title"][:24]))
                else:
                    md5_seen[rec["md5"]] = rec["title"][:24]
        elif media:
            errors.append(f"行{r} 未下载却填了 local_path：{media}")

        # 转写文稿 + 时长（funasr transcript.json 顶层含 duration）
        if media:
            stem = Path(media).stem
            tmd = trans_dir / stem / "transcript.md"
            tjson = trans_dir / stem / "transcript.json"
            if is_tr:
                if tmd.exists():
                    rec["transcript_relpath"] = str(tmd.relative_to(repo))
                else:
                    errors.append(f"行{r} 已转写但缺文稿：{tmd}")
                if tjson.exists():
                    try:
                        dur = json.loads(tjson.read_text(encoding="utf-8")).get("duration")
                        if dur:
                            rec["duration_sec"] = f"{float(dur):.1f}"
                    except Exception as e:
                        errors.append(f"行{r} transcript.json 读取失败：{e}")

        rows.append(rec)

    # 物理目录里有 mp4 但台账没登记（防漏登）
    listed = {Path(r["media_relpath"]).name for r in rows if r["media_relpath"]}
    if video_dir.exists():
        for f in video_dir.glob("*.mp4"):
            if f.name not in listed:
                errors.append(f"视频目录存在但台账未登记：{f.name}")

    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)

    # ---- 对账输出 ----
    print(f"xlsx        : {xlsx.relative_to(repo)}")
    print(f"manifest -> : {out.relative_to(repo)}")
    print(f"记录总数    : {len(rows)}")
    print("价值分级    :", dict(Counter(r["value_level"] or "(空)" for r in rows)))
    print("复查评级    :", dict(Counter(r["review_grade"] or "(空)" for r in rows)))
    print("下载        :", sum(r["downloaded"] == "已下载" for r in rows))
    print("转写        :", sum(r["transcribed"] == "已转写" for r in rows))
    print("md5 已录    :", sum(1 for r in rows if r["md5"]))
    print("时长已录    :", sum(1 for r in rows if r["duration_sec"]))
    print("文稿路径已录:", sum(1 for r in rows if r["transcript_relpath"]))
    if dup_md5:
        print(f"[告警] md5 重复 {len(dup_md5)} 组（同文件一稿两发，非脚本错误）：")
        for h, a, b in dup_md5:
            print(f"   {h[:12]}…  《{a}》 == 《{b}》")
    if errors:
        print("\n[对账失败] 异常 %d 项：" % len(errors))
        for e in errors[:30]:
            print("  -", e)
        sys.exit(1)
    print("\n[对账通过] 状态、文件、指纹、文稿一致。")


if __name__ == "__main__":
    main()
