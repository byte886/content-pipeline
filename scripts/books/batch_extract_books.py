#!/usr/bin/env python3
"""
批量书籍转码脚本：PDF/PPTX → Markdown
- 有文本层的PDF：PyMuPDF直接提取
- 扫描件PDF：Vision OCR（调用processing/ocr/ocr_vision二进制）
- PPTX：python-pptx提取
用法: python3 scripts/books/batch_extract_books.py <输入目录> <输出目录>
"""
import os
import sys
import subprocess
from pathlib import Path

try:
    import pymupdf as fitz
except ImportError:
    import fitz

def extract_pdf_text(pdf_path, output_path):
    """提取PDF文本层，返回(成功, 页数, 平均字符数)"""
    try:
        doc = fitz.open(pdf_path)
        total_pages = len(doc)
        total_chars = 0
        pages_text = []
        for i, page in enumerate(doc):
            text = page.get_text()
            total_chars += len(text)
            if text.strip():
                pages_text.append(f"## 第 {i+1} 页\n\n{text.strip()}\n\n---\n")
        doc.close()
        
        text_page_ratio = len(pages_text) / max(total_pages, 1)
        avg_chars = total_chars / max(total_pages, 1)
        
        # 判断标准：有文本的页数占比>20% 或 平均字符>30，认为有文本层
        if text_page_ratio < 0.2 and avg_chars < 30:
            return False, total_pages, f"文本页占比{text_page_ratio:.0%}, 平均{avg_chars:.0f}字符"
        
        # 写入Markdown
        header = f"# {Path(pdf_path).stem}\n\n> 来源: {Path(pdf_path).name}\n> 页数: {total_pages}\n> 提取方式: PyMuPDF文本层\n> 处理时间: {__import__('datetime').datetime.now().isoformat()}\n\n---\n\n"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(header)
            f.write('\n'.join(pages_text))
        return True, total_pages, avg_chars
    except Exception as e:
        return False, 0, str(e)

def extract_pdf_ocr(pdf_path, output_path, ocr_binary):
    """扫描件PDF：逐页渲染为图片后OCR"""
    try:
        doc = fitz.open(pdf_path)
        total_pages = len(doc)
        pages_text = []
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            for i, page in enumerate(doc):
                # 渲染为图片（200 DPI）
                pix = page.get_pixmap(dpi=200)
                img_path = os.path.join(tmpdir, f"page_{i:04d}.png")
                pix.save(img_path)
                
                # OCR
                result = subprocess.run(
                    [ocr_binary, img_path],
                    capture_output=True, text=True, timeout=60
                )
                text = result.stdout.strip()
                if text:
                    pages_text.append(f"## 第 {i+1} 页\n\n{text}\n\n---\n")
                print(f"  OCR进度: {i+1}/{total_pages}", flush=True)
        doc.close()
        
        header = f"# {Path(pdf_path).stem}\n\n> 来源: {Path(pdf_path).name}\n> 页数: {total_pages}\n> 提取方式: Vision OCR（扫描件）\n> 处理时间: {__import__('datetime').datetime.now().isoformat()}\n\n---\n\n"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(header)
            f.write('\n'.join(pages_text))
        return True, total_pages, "OCR"
    except Exception as e:
        return False, 0, str(e)

def extract_pptx(pptx_path, output_path):
    """提取PPTX文本"""
    try:
        from pptx import Presentation
        prs = Presentation(pptx_path)
        slides_text = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides_text.append(f"## 第 {i+1} 页\n\n" + "\n\n".join(texts) + "\n\n---\n")
        
        header = f"# {Path(pptx_path).stem}\n\n> 来源: {Path(pptx_path).name}\n> 页数: {len(prs.slides)}\n> 提取方式: python-pptx\n> 处理时间: {__import__('datetime').datetime.now().isoformat()}\n\n---\n\n"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(header)
            f.write('\n'.join(slides_text))
        return True, len(prs.slides), "PPTX"
    except Exception as e:
        return False, 0, str(e)

def main():
    if len(sys.argv) < 3:
        print("用法: python3 batch_extract_books.py <输入目录> <输出目录>")
        sys.exit(1)
    
    input_dir = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])
    output_dir.mkdir(parents=True, exist_ok=True)
    
    ocr_binary = "processing/ocr/ocr_vision"
    if not os.path.exists(ocr_binary):
        ocr_binary = None
        print("⚠️  未找到OCR二进制，扫描件将跳过")
    
    files = sorted(list(input_dir.glob("*.pdf")) + list(input_dir.glob("*.pptx")))
    print(f"找到 {len(files)} 个文件")
    print("=" * 60)
    
    results = []
    for idx, f in enumerate(files, 1):
        print(f"\n[{idx}/{len(files)}] 处理: {f.name}")
        output_path = output_dir / f"{f.stem}_extracted.md"
        
        if output_path.exists():
            print(f"  已存在，跳过: {output_path.name}")
            results.append((f.name, "跳过", ""))
            continue
        
        if f.suffix.lower() == '.pdf':
            success, pages, info = extract_pdf_text(str(f), str(output_path))
            if success:
                print(f"  ✅ 文本层提取: {pages}页, 平均{info:.0f}字符/页")
                results.append((f.name, "文本层", f"{pages}页"))
            else:
                print(f"  ⚠️  文本层不足({info})，尝试OCR...")
                if ocr_binary:
                    success2, pages2, info2 = extract_pdf_ocr(str(f), str(output_path), ocr_binary)
                    if success2:
                        print(f"  ✅ OCR提取: {pages2}页")
                        results.append((f.name, "OCR", f"{pages2}页"))
                    else:
                        print(f"  ❌ OCR失败: {info2}")
                        results.append((f.name, "失败", str(info2)))
                else:
                    results.append((f.name, "需OCR", "无OCR工具"))
        elif f.suffix.lower() == '.pptx':
            success, pages, info = extract_pptx(str(f), str(output_path))
            if success:
                print(f"  ✅ PPTX提取: {pages}页")
                results.append((f.name, "PPTX", f"{pages}页"))
            else:
                print(f"  ❌ PPTX失败: {info}")
                results.append((f.name, "失败", str(info)))
    
    print("\n" + "=" * 60)
    print("处理汇总:")
    for name, method, info in results:
        print(f"  [{method}] {name} - {info}")
    
    # 写入汇总
    summary_path = output_dir / "_extract_summary.md"
    with open(summary_path, 'w', encoding='utf-8') as f:
        f.write("# 书籍转码汇总\n\n")
        f.write(f"| 文件名 | 提取方式 | 信息 |\n|---|---|---|\n")
        for name, method, info in results:
            f.write(f"| {name} | {method} | {info} |\n")
    print(f"\n汇总已写入: {summary_path}")

if __name__ == "__main__":
    main()
